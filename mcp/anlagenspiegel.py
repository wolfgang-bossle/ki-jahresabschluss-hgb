"""
anlagenspiegel.py
Zweck: Engine — Anlagenbuchhaltung (Brutto) + Tabelle B + XBRL-Taxonomie
       => Anlagenspiegel (§284 Abs. 3 / §268 Abs. 2 HGB) als Datenmodell-Block.
       Harte Reconciliation gegen die Saldenliste (BW Ende = Anlagekonten,
       Σ AfA Jahr = AfA-Aufwandskonto) — Abweichung = ValueError, nie still.
Status: ✅ Produktiv 2026-06-18
Abhängigkeiten: data/<mandant>/Anlagenbuchhaltung.xlsx, mcp/config/skr03_mapping.json,
                taxonomy/*-presentation-balanceSheet.xml, *-label-de.xml, openpyxl
Datenagnostisch (§2.1): KEINE Kontonummern/Mandant im Code. Werte aus der
                Anlagenbuchhaltung, Zuordnung aus Tabelle B, Gliederung aus der
                Taxonomie (§2.8). Bilanz/GuV werden NICHT angefasst (additiver Block).
Letzte Änderung: 2026-06-18
"""
import json
from collections import defaultdict
from pathlib import Path

import openpyxl

# Reuse der geprüften Bausteine aus der Bilanz/GuV-Engine (kein Duplikat)
from jahresabschluss import build_tree, load_labels, load_mapping, read_saldenliste

# --- Taxonomie-Strukturkonstanten (NICHT mandantenspezifisch, aus de-gaap-ci) ---
# Anlagenspiegel deckt das Anlagevermögen ab; bei der Bäckerei nur Sachanlagen.
ANLAGEVERMOEGEN_ROOT = "de-gaap-ci_bs.ass.fixAss"
# Marker des AfA-Aufwandskonzepts in der GuV (für den AfA-Abgleich, ohne Kontonummer):
DEPR_MARKER = ".deprAmort"

# Bewegungsspalten der Brutto-Anlagenbuchhaltung (Reihenfolge = Spalten der xlsx ab Index 2)
COLS = ["ahk_anf", "zugang", "abgang_ahk", "umbuchung", "ahk_ende",
        "kumafa_anf", "afa_jahr", "abgang_afa", "kumafa_ende", "bw_gj", "bw_vj"]
_COLIDX = {c: i for i, c in enumerate(COLS, start=2)}  # ahk_anf=Spalte 2 ... bw_vj=Spalte 12


# --------------------------------------------------------------------------- #
# Mandantendaten laden
# --------------------------------------------------------------------------- #
def read_anlagenbuchhaltung(xlsx_path):
    """Konto -> {bewegungsspalten..., bezeichnung, methode, nd}. Brutto-Format."""
    ws = openpyxl.load_workbook(xlsx_path, data_only=True)["Anlagenbuchhaltung"]
    out = {}
    for row in ws.iter_rows(values_only=True):
        if not row[0] or not str(row[0])[0].isdigit():
            continue
        konto = str(row[0]).strip()
        rec = {c: float(row[i] or 0) for c, i in _COLIDX.items()}
        rec["bezeichnung"] = str(row[1] or "")
        rec["methode"] = str(row[13] or "")
        rec["nd"] = str(row[14] or "")
        out[konto] = rec
    return out


def read_hinweise(xlsx_path):
    """Optionales Blatt 'Hinweise' -> [{gruppe, hinweis, typ}]."""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    if "Hinweise" not in wb.sheetnames:
        return []
    out = []
    for row in wb["Hinweise"].iter_rows(min_row=3, values_only=True):
        if row and row[0] and row[1]:
            out.append({"gruppe": str(row[0]), "hinweis": str(row[1]),
                        "typ": str(row[2] or "")})
    return out


# --------------------------------------------------------------------------- #
# Engine
# --------------------------------------------------------------------------- #
def generate_anlagenspiegel(anlagenbuchhaltung, saldenliste, mapping_file, taxonomy_dir):
    tax = Path(taxonomy_dir)
    labels = load_labels(tax / "de-gaap-ci-2025-04-01-label-de.xml")
    anlagen = read_anlagenbuchhaltung(anlagenbuchhaltung)
    hinweise = read_hinweise(anlagenbuchhaltung)
    konto2concept = load_mapping(mapping_file)
    konten_sl = read_saldenliste(saldenliste)

    # (1) Jede Anlage muss eine Tabelle-B-Zuordnung haben
    fehlend = [k for k in anlagen if k not in konto2concept]
    if fehlend:
        raise ValueError(f"Anlagen-Konten ohne Tabelle-B-Zuordnung: {fehlend}")

    # (2) Harte Reconciliation gegen die Saldenliste (eiserner Grundsatz §2.7)
    for konto, rec in anlagen.items():
        sl = konten_sl.get(konto)
        if sl is None:
            raise ValueError(f"Anlagen-Konto {konto} fehlt in der Saldenliste")
        _, sgj, svj = sl
        if round(rec["bw_gj"], 2) != round(sgj, 2) or round(rec["bw_vj"], 2) != round(svj, 2):
            raise ValueError(
                f"BW-Abweichung Konto {konto}: Anlagenbuch {rec['bw_gj']}/{rec['bw_vj']} "
                f"!= Saldenliste {sgj}/{svj} -> Anlagenbuchhaltung stimmt nicht mit Saldenliste überein")
        if round(rec["ahk_ende"] - rec["kumafa_ende"], 2) != round(rec["bw_gj"], 2):
            raise ValueError(f"Zeile {konto}: AHK Ende − Kum.AfA Ende != BW GJ")
        if round(rec["ahk_anf"] - rec["kumafa_anf"], 2) != round(rec["bw_vj"], 2):
            raise ValueError(f"Zeile {konto}: AHK Anf − Kum.AfA Anf != BW VJ")

    # Σ AfA Jahr (Anlagenspiegel) == AfA-Aufwand laut Saldenliste (Konzept ...deprAmort)
    afa_konten = [k for k, c in konto2concept.items() if DEPR_MARKER in c and k in konten_sl]
    afa_soll = round(sum(konten_sl[k][1] for k in afa_konten), 2)
    afa_ist = round(sum(rec["afa_jahr"] for rec in anlagen.values()), 2)
    if afa_ist != afa_soll:
        raise ValueError(
            f"Σ AfA Anlagenspiegel {afa_ist} != AfA-Aufwand Saldenliste {afa_soll} "
            f"(Konten {afa_konten}) -> Abschreibung nicht abgestimmt")

    # (3) Leaf-Aggregation je Taxonomie-Konzept + beitragende Konten
    concept_vals = defaultdict(lambda: {c: 0.0 for c in COLS})
    concept_konten = defaultdict(list)
    for konto, rec in anlagen.items():
        c = konto2concept[konto]
        for col in COLS:
            concept_vals[c][col] += rec[col]
        concept_konten[c].append(konto)

    # (4) Rollup entlang der Bilanz-Presentation-Hierarchie (Konzept -> §266 A.II)
    c2p, p2c = build_tree(tax / "de-gaap-ci-2025-04-01-presentation-balanceSheet.xml",
                          set(concept_vals))
    node = defaultdict(lambda: {c: 0.0 for c in COLS})
    node_konten = defaultdict(set)
    for c, vals in concept_vals.items():
        cur, seen = c, set()
        while cur and cur not in seen:
            seen.add(cur)
            for col in COLS:
                node[cur][col] += vals[col]
            node_konten[cur].update(concept_konten[c])
            cur = c2p.get(cur)

    # (5) Geordnete, flache Positionsliste ab dem Anlagevermögen; je Blatt-Konzept
    #     hängen die Einzelkonten als unterste Ebene (volle Rückverfolgbarkeit).
    rows = []

    def walk(c, depth):
        if c not in node:
            return
        rows.append({
            "ebene": depth, "konzept": c, "konto": None,
            "label": labels.get(c, c), "quelle": sorted(node_konten[c]),
            **{col: round(node[c][col], 2) for col in COLS},
        })
        for _, ch in p2c.get(c, []):
            if ch in node:
                walk(ch, depth + 1)
        for konto in sorted(concept_konten.get(c, [])):
            rec = anlagen[konto]
            rows.append({
                "ebene": depth + 1, "konzept": None, "konto": konto,
                "label": rec["bezeichnung"], "quelle": [konto],
                "methode": rec["methode"], "nd": rec["nd"],
                **{col: round(rec[col], 2) for col in COLS},
            })

    walk(ANLAGEVERMOEGEN_ROOT, 0)
    summe = {col: round(node[ANLAGEVERMOEGEN_ROOT][col], 2) for col in COLS}

    return {
        "positionen": rows,
        "summe": summe,
        "hinweise": hinweise,
        "reconciliation": {
            "bw_gj": summe["bw_gj"], "bw_vj": summe["bw_vj"],
            "afa_jahr": afa_ist, "afa_aufwand_saldenliste": afa_soll,
            "afa_konten": sorted(afa_konten), "abgestimmt": True,
        },
    }


if __name__ == "__main__":
    import sys
    base = Path(__file__).resolve().parent.parent
    asp = generate_anlagenspiegel(
        base / "data/baeckerei_2025/Anlagenbuchhaltung.xlsx",
        base / "data/baeckerei_2025/Saldenliste.xlsx",
        base / "mcp/config/skr03_mapping.json",
        base / "taxonomy",
    )
    sys.stdout.reconfigure(encoding="utf-8")
    print("=" * 110, "\nANLAGENSPIEGEL (§284 Abs. 3 HGB) — Brutto")
    hdr = f"{'Position':<40}{'AHK Anf':>12}{'Zugang':>11}{'AHK Ende':>12}{'AfA Jahr':>11}{'BW GJ':>12}{'BW VJ':>12}"
    print(hdr); print("-" * 110)
    for r in asp["positionen"]:
        ind = "  " * r["ebene"]
        name = (ind + r["label"])[:40]
        print(f"{name:<40}{r['ahk_anf']:>12,.0f}{r['zugang']:>11,.0f}{r['ahk_ende']:>12,.0f}"
              f"{r['afa_jahr']:>11,.0f}{r['bw_gj']:>12,.0f}{r['bw_vj']:>12,.0f}")
    s = asp["summe"]
    print("-" * 110)
    print(f"{'SUMME ANLAGEVERMÖGEN':<40}{s['ahk_anf']:>12,.0f}{s['zugang']:>11,.0f}"
          f"{s['ahk_ende']:>12,.0f}{s['afa_jahr']:>11,.0f}{s['bw_gj']:>12,.0f}{s['bw_vj']:>12,.0f}")
    rc = asp["reconciliation"]
    print(f"\nReconciliation: BW GJ {rc['bw_gj']:,.0f} | BW VJ {rc['bw_vj']:,.0f} | "
          f"AfA {rc['afa_jahr']:,.0f} = Saldenliste-AfA {rc['afa_aufwand_saldenliste']:,.0f} "
          f"(Konten {rc['afa_konten']}) | abgestimmt={rc['abgestimmt']}")
    for h in asp["hinweise"]:
        print(f"Hinweis [{h['typ']}] {h['gruppe']}: {h['hinweis']}")
